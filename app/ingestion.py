"""Bounded CSV/XLSX parsing and explicit, reproducible field normalization."""
import csv
import hashlib
import io
import ipaddress
import json
import re
import unicodedata
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .config import MAX_ROWS, MAX_UPLOAD, TIMEZONE

FIELDS = {
    'timestamp': ['timestamp', 'time', 'datetime', 'fecha', 'fechahora', 'date', 'timecreated', 'eventtime', '@timestamp'],
    'src_ip': ['src_ip', 'sourceip', 'source_ip', 'iporigen', 'srcip', 'ipaddress', 'clientip', 'sourceaddress'],
    'dst_ip': ['dst_ip', 'destinationip', 'destination_ip', 'ipdestino', 'dstip', 'destinationaddress'],
    'username': ['username', 'user', 'usuario', 'accountname', 'targetusername', 'user_name'],
    'hostname': ['hostname', 'host', 'equipo', 'computer', 'computername', 'device', 'desthost'],
    'event_type': ['event_type', 'eventtype', 'tipo', 'tipoevento', 'category', 'action', 'actividad'],
    'event_id': ['event_id', 'eventid', 'idevento', 'id', 'eventcode'],
    'status': ['status', 'result', 'resultado', 'estado', 'outcome'],
    'dst_port': ['dst_port', 'destinationport', 'destination_port', 'dstport', 'puertodestino', 'port'],
    'bytes_sent': ['bytes_sent', 'bytessent', 'bytesout', 'sentbytes'],
    'message': ['message', 'mensaje', 'description', 'descripcion', 'commandline', 'command', 'scriptblocktext', 'details'],
}
LEGACY_FIELDS = tuple(FIELDS)
FIELDS.update({
    'log_id': ['Log_ID', 'logid', 'source_record_id'],
    'destination': ['Destination', 'destino'],
    'log_source': ['Log_Source', 'logsource', 'fuente'],
    'event_description': ['Event_Description', 'eventdescription'],
    'process_commandline': ['Process_CommandLine', 'processcommandline'],
    'source_severity': ['Severity', 'source_severity', 'severidad'],
})


def key(value):
    value = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]', '', value)


class SourceRow(list):
    """Preserve the physical source row even when blank rows are skipped."""
    def __init__(self, values, source_row, sheet_name=None):
        super().__init__(values)
        self.source_row = source_row
        self.sheet_name = sheet_name


def read_table(filename, content):
    if not content or len(content) > MAX_UPLOAD:
        raise ValueError('El archivo debe contener datos y pesar como máximo 10 MB.')
    suffix = Path(filename).suffix.lower()
    if suffix == '.csv':
        for encoding in ('utf-8-sig', 'cp1252'):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError('Codificación no compatible. Guarda el CSV como UTF-8.')
        if '\x00' in decoded:
            raise ValueError('El CSV contiene datos binarios. Utiliza UTF-8 o Windows-1252.')
        try:
            dialect = csv.Sniffer().sniff(decoded[:16000], delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel
        table = csv.reader(io.StringIO(decoded, newline=''), dialect)
        rows = []
        previous_line = 0
        try:
            for row in table:
                first_line = previous_line + 1
                previous_line = table.line_num
                if any(str(cell).strip() for cell in row):
                    rows.append(SourceRow(row, first_line))
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError(f'Máximo {MAX_ROWS:,} filas por archivo.')
        except csv.Error as exc:
            raise ValueError('El CSV está mal formado o contiene un campo demasiado grande.') from exc
    elif suffix == '.xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(x.file_size for x in archive.infolist()) > 60 * 1024 * 1024:
                    raise ValueError('El Excel supera el límite de 60 MB descomprimidos o 2.000 entradas.')
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.worksheets[0]
                if sheet.max_column and sheet.max_column > 100:
                    raise ValueError('Máximo 100 columnas por hoja.')
                rows = []
                for source_row, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if any(cell is not None for cell in row):
                        if any(isinstance(cell, str) and cell.startswith('=') for cell in row):
                            raise ValueError('El Excel contiene fórmulas. Exporta los valores a un archivo CSV.')
                        rows.append(SourceRow(row, source_row, sheet.title))
                    if len(rows) > MAX_ROWS + 1:
                        raise ValueError(f'Máximo {MAX_ROWS:,} filas por archivo.')
            finally:
                workbook.close()
        except (zipfile.BadZipFile, KeyError, OSError, ParseError, InvalidFileException, EOFError) as exc:
            raise ValueError('El archivo no es un Excel .xlsx válido.') from exc
    else:
        raise ValueError('Formatos admitidos: .csv y .xlsx. Convierte los archivos .xls a .xlsx.')
    if len(rows) < 2:
        raise ValueError('Se requiere una fila de encabezados y al menos una fila de datos.')
    headers = [str(cell or '').strip() for cell in rows[0]]
    if len(headers) > 100 or len(set(headers)) != len(headers) or not all(headers):
        raise ValueError('Utiliza entre 1 y 100 encabezados únicos, sin celdas vacías.')
    return headers, rows[1:]


def suggest_mapping(headers):
    mapping = {}
    for field, aliases in FIELDS.items():
        candidates = {key(alias) for alias in aliases + [field]}
        found = next((h for h in headers if key(h) in candidates), None)
        if found:
            mapping[field] = found
    return mapping


def parse_time(value, source_timezone=TIMEZONE):
    """Legacy callers retain their default; the upload API explicitly passes None.

    None preserves a naive timestamp instead of inventing an offset. Only an
    embedded offset or an explicitly supplied zone allows UTC conversion.
    """
    if isinstance(value, datetime):
        parsed = value
    else:
        text = str(value or '').strip()
        if not text:
            raise ValueError('Fecha ausente.')
        try:
            parsed = datetime.fromisoformat(text.replace('Z', '+00:00'))
        except ValueError:
            parsed = None
            for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%Y/%m/%d %H:%M:%S'):
                try:
                    parsed = datetime.strptime(text, fmt)
                    break
                except ValueError:
                    pass
            if parsed is None:
                raise ValueError('Fecha inválida. Usa ISO 8601 o DD/MM/AAAA HH:MM:SS.')
    if parsed.tzinfo is None:
        if source_timezone is None:
            return parsed.isoformat()
        parsed = parsed.replace(tzinfo=ZoneInfo(source_timezone))
    return parsed.astimezone(timezone.utc).isoformat()


def normalize(headers, rows, mapping=None, source_timezone=TIMEZONE):
    mapping = suggest_mapping(headers) if mapping is None else mapping
    if not isinstance(mapping, dict) or any(f not in FIELDS or h not in headers for f, h in mapping.items()):
        raise ValueError('El mapeo contiene campos o columnas no válidos.')
    if 'timestamp' not in mapping:
        raise ValueError('Asigna una columna a Fecha / hora antes de importar.')
    if not any(f in mapping for f in ('event_type', 'event_id', 'status', 'message')):
        raise ValueError('Asigna al menos un campo de contexto: tipo, ID de evento, resultado o mensaje.')
    events, errors = [], []
    positions = {field: headers.index(column) for field, column in mapping.items()}
    for number, row in enumerate(rows, start=2):
        number = getattr(row, 'source_row', number)
        try:
            if len(row) != len(headers):
                raise ValueError('La cantidad de columnas no coincide con el encabezado.')
            event = {f: '' for f in FIELDS}
            for field, index in positions.items():
                value = row[index]
                if isinstance(value, float) and value.is_integer():
                    value = int(value)
                event[field] = str(value if value is not None else '').strip()
            original_time = row[positions['timestamp']]
            event['timestamp'] = parse_time(original_time, source_timezone)
            parsed_time = datetime.fromisoformat(event['timestamp'])
            # Destination is an opaque source value: hosts and URLs are not invalid IPs.
            if event['destination'] and not event['dst_ip']:
                try:
                    event['dst_ip'] = str(ipaddress.ip_address(event['destination']))
                except ValueError:
                    pass
            if not event['message']:
                event['message'] = '\n'.join(v for v in (event['event_description'], event['process_commandline']) if v)
            for field in ('src_ip', 'dst_ip'):
                if event[field] and event[field].lower() not in ('-', '::', '0.0.0.0', 'n/a', 'na', 'null', 'none', 'unknown'):
                    event[field] = str(ipaddress.ip_address(event[field]))
                else:
                    event[field] = ''
            for field in ('dst_port', 'bytes_sent'):
                if event[field]:
                    value = int(event[field])
                    if value < 0 or (field == 'dst_port' and value > 65535):
                        raise ValueError(f'{field}: valor fuera de rango.')
                    event[field] = value
                else:
                    event[field] = 0
            event['event_id'] = event['event_id'].removesuffix('.0')
            event['username'] = event['username'].lower()
            event['hostname'] = event['hostname'].lower()
            if any(len(str(v)) > 12000 for v in event.values()):
                raise ValueError('Un campo supera el límite de 12.000 caracteres.')
            raw = {str(h): str(v) if v is not None else '' for h, v in zip(headers, row)}
            # Empty new fields must not change fingerprints of old, offset-aware logs.
            identity = {field: event[field] for field in LEGACY_FIELDS}
            identity.update({field: event[field] for field in FIELDS if field not in LEGACY_FIELDS and event[field]})
            event['fingerprint'] = hashlib.sha256(json.dumps(identity, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
            event['timestamp_original'] = str(original_time)
            event['timestamp_utc'] = event['timestamp'] if parsed_time.tzinfo else None
            event['timezone_status'] = 'known' if parsed_time.tzinfo else 'unknown'
            event['source_timezone'] = source_timezone
            event['raw'] = raw
            event['row_number'] = number
            event['sheet_name'] = getattr(row, 'sheet_name', None)
            events.append(event)
        except (ValueError, TypeError, OverflowError) as exc:
            errors.append({'row': number, 'error': str(exc),
                           'log_id': str(row[positions['log_id']]) if 'log_id' in positions and len(row) > positions['log_id'] else ''})
    return events, errors, mapping
